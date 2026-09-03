import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\vega_\Downloads\20260830-兴业证券-化工行业：化工品价格价差与库存开工数据库.xlsx', data_only=True, read_only=True)
ws = wb['涨跌幅更新']
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 4:
        print(f"Row {i}:", [str(c) for c in row[:18]])

print("---")
wb2 = openpyxl.load_workbook(r'C:\Users\vega_\Downloads\申万宏源-化工细分行业及重点公司一周概览20260828.xlsx', data_only=True, read_only=True)
ws2 = wb2['价格变化总表']
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i < 5:
        print(f"SW Row {i}:", [str(c) for c in row[:15]])
